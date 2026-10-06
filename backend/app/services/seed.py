"""Load reference data (CSV/XLSX under data/) and demo users into the database."""
from __future__ import annotations

import csv
import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.config import get_settings
from ..core.security import hash_password
from ..models import (HistoricalQuote, ManufacturingCost, MarginRule, Product, ShippingRate,
                      SpecParameter, Supplier, SupplierPrice, User)


def _read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _read_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]


def _f(v) -> float | None:
    return None if v in ("", None) else float(v)


def _ref(rel: str, row_id) -> str:
    return f"{rel}#id={row_id}"


def reference_data_loaded(db: Session) -> bool:
    return db.scalar(select(Product.id).limit(1)) is not None


def seed_reference_data(db: Session, data_dir: Path | None = None) -> dict:
    root = data_dir or get_settings().data_dir
    if reference_data_loaded(db):
        return {"status": "skipped"}
    p_rows = _read_csv(root / "products" / "products.csv")
    for r in p_rows:
        db.add(Product(id=int(r["id"]), code=r["code"], name=r["name"], category=r["category"],
                       hazmat_class=r["hazmat_class"] or None, hs_code=r["hs_code"], description=r["description"],
                       source_ref=_ref("products/products.csv", r["id"])))
    db.flush()
    pid = {r["code"]: int(r["id"]) for r in p_rows}

    for r in _read_csv(root / "specifications" / "specs.csv"):
        db.add(SpecParameter(id=int(r["id"]), product_id=pid[r["product_code"]], parameter=r["parameter"],
                             operator=r["operator"], min_value=_f(r["min_value"]), max_value=_f(r["max_value"]),
                             unit=r["unit"], critical=bool(int(r["critical"])),
                             review_margin=float(r["review_margin"]),
                             source_ref=_ref("specifications/specs.csv", r["id"])))
    s_rows = _read_csv(root / "suppliers" / "suppliers.csv")
    for r in s_rows:
        db.add(Supplier(id=int(r["id"]), code=r["code"], name=r["name"], country=r["country"],
                        lead_time_days=int(r["lead_time_days"]), moq_kg=float(r["moq_kg"]),
                        rating=float(r["rating"]), source_ref=_ref("suppliers/suppliers.csv", r["id"])))
    db.flush()
    sid = {r["code"]: int(r["id"]) for r in s_rows}

    for r in _read_csv(root / "pricing" / "supplier_prices.csv"):
        db.add(SupplierPrice(id=int(r["id"]), supplier_id=sid[r["supplier_code"]], product_id=pid[r["product_code"]],
                             purity_grade=float(r["purity_grade"]), tier_min_kg=float(r["tier_min_kg"]),
                             price_per_kg=Decimal(r["price_per_kg"]), currency=r["currency"],
                             valid_until=dt.date.fromisoformat(r["valid_until"]),
                             source_ref=_ref("pricing/supplier_prices.csv", r["id"])))
    for r in _read_csv(root / "pricing" / "manufacturing_costs.csv"):
        db.add(ManufacturingCost(id=int(r["id"]), product_id=pid[r["product_code"]], plant_name=r["plant_name"],
                                 country=r["country"], purity_grade=float(r["purity_grade"]),
                                 cost_per_kg=Decimal(r["cost_per_kg"]), min_batch_kg=float(r["min_batch_kg"]),
                                 capacity_kg_month=float(r["capacity_kg_month"]),
                                 lead_time_days=int(r["lead_time_days"]),
                                 source_ref=_ref("pricing/manufacturing_costs.csv", r["id"])))
    for r in _read_xlsx(root / "pricing" / "margin_rules.xlsx"):
        db.add(MarginRule(id=int(r["id"]), category=r["category"], qty_tier_min_kg=float(r["qty_tier_min_kg"]),
                          customer_tier=r["customer_tier"], margin_pct=Decimal(str(r["margin_pct"])),
                          source_ref=_ref("pricing/margin_rules.xlsx", r["id"])))
    for r in _read_csv(root / "shipping" / "shipping_rates.csv"):
        db.add(ShippingRate(id=int(r["id"]), origin_country=r["origin_country"], dest_port=r["dest_port"],
                            dest_country=r["dest_country"], mode=r["mode"], rate_per_kg=Decimal(r["rate_per_kg"]),
                            min_charge=Decimal(r["min_charge"]), transit_days=int(r["transit_days"]),
                            hazmat_surcharge_pct=Decimal(r["hazmat_surcharge_pct"]),
                            source_ref=_ref("shipping/shipping_rates.csv", r["id"])))
    for r in _read_csv(root / "historical_quotes" / "historical_quotes.csv"):
        db.add(HistoricalQuote(id=int(r["id"]), product_id=pid[r["product_code"]], qty_kg=float(r["qty_kg"]),
                               dest_port=r["dest_port"], quoted_price_per_kg=Decimal(r["quoted_price_per_kg"]),
                               outcome=r["outcome"], quoted_at=dt.date.fromisoformat(r["quoted_at"]),
                               source_ref=_ref("historical_quotes/historical_quotes.csv", r["id"])))
    db.commit()
    return {"status": "seeded", "products": len(p_rows), "suppliers": len(s_rows)}


DEMO_USERS = [("analyst@opspilot.demo", "Asha Analyst", "analyst"),
              ("approver@opspilot.demo", "Arjun Approver", "approver"),
              ("admin@opspilot.demo", "Admin", "admin")]


def seed_users(db: Session) -> int:
    pw = get_settings().demo_password
    n = 0
    for email, name, role in DEMO_USERS:
        if db.scalar(select(User).where(User.email == email)) is None:
            db.add(User(email=email, display_name=name, role=role, password_hash=hash_password(pw)))
            n += 1
    db.commit()
    return n
