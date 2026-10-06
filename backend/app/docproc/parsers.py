"""Parse PDF / TXT / CSV / XLSX bytes into text segments with locators."""
from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field

from openpyxl import load_workbook
from pypdf import PdfReader


class ParseError(Exception):
    pass


@dataclass
class Segment:
    text: str
    locator: str


@dataclass
class ParsedDocument:
    filename: str
    kind: str  # pdf | txt | csv | xlsx
    segments: list[Segment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(s.text for s in self.segments)


def parse_bytes(filename: str, content: bytes, max_pdf_pages: int = 30,
                max_rows: int = 5000) -> ParsedDocument:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _parse_pdf(filename, content, max_pdf_pages)
    if ext == ".txt":
        return _parse_txt(filename, content)
    if ext == ".csv":
        return _parse_csv(filename, content, max_rows)
    if ext == ".xlsx":
        return _parse_xlsx(filename, content, max_rows)
    raise ParseError(f"Unsupported file type: {ext}")


def _parse_pdf(filename: str, content: bytes, max_pages: int) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            if not reader.decrypt(""):
                raise ParseError("PDF is password protected")
        doc = ParsedDocument(filename, "pdf")
        pages = reader.pages
        if len(pages) > max_pages:
            doc.warnings.append(f"Only the first {max_pages} of {len(pages)} pages were read")
        for i, page in enumerate(pages[:max_pages], start=1):
            text = (page.extract_text() or "").strip()
            if text:
                doc.segments.append(Segment(text, f"page {i}"))
    except ParseError:
        raise
    except Exception as exc:  # pypdf raises many exception types on malformed input
        raise ParseError(f"Could not read PDF: {exc.__class__.__name__}") from None
    if not doc.segments:
        raise ParseError("PDF has no extractable text (scanned image?). OCR is not supported in this prototype.")
    return doc


def _parse_txt(filename: str, content: bytes) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace").replace("\r\n", "\n").strip()
    if not text:
        raise ParseError("Text file is empty")
    return ParsedDocument(filename, "txt", [Segment(text, "text")])


def _rows_to_segments(rows: list[list[str]], prefix: str = "") -> list[Segment]:
    if not rows:
        return []
    header = [h.strip() for h in rows[0]]
    segs = []
    for n, row in enumerate(rows[1:], start=2):
        if not any(str(c).strip() for c in row):
            continue
        pairs = [f"{h}: {str(v).strip()}" for h, v in zip(header, row) if str(v).strip() != ""]
        segs.append(Segment(prefix + "; ".join(pairs), f"row {n}"))
    return segs


def _parse_csv(filename: str, content: bytes, max_rows: int) -> ParsedDocument:
    text = content.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    doc = ParsedDocument(filename, "csv")
    if len(rows) > max_rows + 1:
        doc.warnings.append(f"Only the first {max_rows} rows were read")
        rows = rows[: max_rows + 1]
    doc.segments = _rows_to_segments(rows)
    if not doc.segments:
        raise ParseError("CSV has no data rows")
    return doc


def _parse_xlsx(filename: str, content: bytes, max_rows: int) -> ParsedDocument:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Could not read workbook: {exc.__class__.__name__}") from None
    doc = ParsedDocument(filename, "xlsx")
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > max_rows:
                doc.warnings.append(f"Sheet '{ws.title}' truncated at {max_rows} rows")
                break
            rows.append(["" if c is None else str(c) for c in row])
        for seg in _rows_to_segments(rows, prefix=f"[{ws.title}] "):
            doc.segments.append(Segment(seg.text, f"{ws.title} {seg.locator}"))
    wb.close()
    if not doc.segments:
        raise ParseError("Workbook has no data rows")
    return doc
